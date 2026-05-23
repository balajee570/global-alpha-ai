"""
NASDAQ Excel Export - Multi-sheet workbook builder
Exports comprehensive analysis to Excel with formatting
"""

import pandas as pd
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _get_header_style():
    """Return header cell style."""
    return {
        "fill": PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid"),
        "font": Font(color="FFFFFF", bold=True, size=11),
        "alignment": Alignment(horizontal="center", vertical="center"),
        "border": Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )
    }


def _get_cell_style():
    """Return standard cell style."""
    return {
        "border": Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        ),
        "alignment": Alignment(horizontal="left", vertical="top", wrap_text=True)
    }


def _add_dataframe_to_sheet(ws, df, start_row=1, header_style=None, cell_style=None):
    """Add DataFrame to worksheet with formatting."""
    if df.empty:
        return

    header_style = header_style or _get_header_style()
    cell_style = cell_style or _get_cell_style()

    # Write header
    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=col_name)
        for key, value in header_style.items():
            setattr(cell, key, value)

    # Write data
    for row_idx, (_, row) in enumerate(df.iterrows(), start_row + 1):
        for col_idx, value in enumerate(row.values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            for key, value_style in cell_style.items():
                setattr(cell, key, value_style)

    # Auto-adjust column widths
    for col_idx, col_name in enumerate(df.columns, 1):
        max_length = max(
            df[col_name].astype(str).map(len).max(),
            len(str(col_name))
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 50)


def build_excel_workbook(scan_result: dict, fund_map: dict) -> io.BytesIO:
    """
    Build comprehensive Excel workbook with 6+ sheets.

    Returns BytesIO buffer ready for download.
    """
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    momentum_df = scan_result.get("momentum_df", pd.DataFrame())
    shortlist_df = scan_result.get("shortlist_df", pd.DataFrame())
    regime = scan_result.get("regime", {})
    sector_strength = scan_result.get("sector_strength", pd.DataFrame())

    # ========== Sheet 1: Full Scan ==========
    ws = wb.create_sheet("Full Scan", 0)
    _add_dataframe_to_sheet(ws, momentum_df)
    ws.freeze_panes = "A2"

    # ========== Sheet 2: Summary (Shortlist) ==========
    ws = wb.create_sheet("Summary", 1)
    if not shortlist_df.empty:
        summary_df = shortlist_df.copy()
        summary_df["Action"] = [
            fund_map.get(t, {}).get("recommendation", {}).get("action", "HOLD")
            for t in summary_df["Ticker"]
        ]
        summary_df["Conviction"] = [
            fund_map.get(t, {}).get("recommendation", {}).get("conviction", 0)
            for t in summary_df["Ticker"]
        ]
        _add_dataframe_to_sheet(ws, summary_df)
    ws.freeze_panes = "A2"

    # ========== Sheet 3: Valuation ==========
    ws = wb.create_sheet("Valuation", 2)
    val_data = []
    for _, row in shortlist_df.iterrows():
        ticker = row["Ticker"]
        metrics = fund_map.get(ticker, {}).get("metrics", {})
        val_data.append({
            "Ticker": ticker,
            "Price": metrics.get("price"),
            "P/E (Trailing)": metrics.get("pe_trailing"),
            "P/E (Forward)": metrics.get("pe_forward"),
            "P/B": metrics.get("pb"),
            "PEG": metrics.get("peg"),
            "EV/EBITDA": metrics.get("ev_ebitda"),
            "EV/FCF": metrics.get("ev_fcf"),
            "P/S": metrics.get("ps"),
            "Div Yield %": metrics.get("dividend_yield"),
        })
    val_df = pd.DataFrame(val_data)
    if not val_df.empty:
        _add_dataframe_to_sheet(ws, val_df)
    ws.freeze_panes = "A2"

    # ========== Sheet 4: Quality ==========
    ws = wb.create_sheet("Quality", 3)
    qual_data = []
    for _, row in shortlist_df.iterrows():
        ticker = row["Ticker"]
        metrics = fund_map.get(ticker, {}).get("metrics", {})
        qual_data.append({
            "Ticker": ticker,
            "Piotroski": metrics.get("piotroski_score"),
            "ROE %": metrics.get("roe"),
            "ROCE %": metrics.get("roce"),
            "ROA %": metrics.get("roa"),
            "Gross Margin %": metrics.get("gross_margin"),
            "Op Margin %": metrics.get("op_margin"),
            "Net Margin %": metrics.get("net_margin"),
            "Margin Trend (3Y)": metrics.get("net_margin_trend"),
            "D/E": metrics.get("debt_to_equity"),
            "Current Ratio": metrics.get("current_ratio"),
            "Interest Cover": metrics.get("interest_coverage"),
        })
    qual_df = pd.DataFrame(qual_data)
    if not qual_df.empty:
        _add_dataframe_to_sheet(ws, qual_df)
    ws.freeze_panes = "A2"

    # ========== Sheet 5: Growth ==========
    ws = wb.create_sheet("Growth", 4)
    growth_data = []
    for _, row in shortlist_df.iterrows():
        ticker = row["Ticker"]
        metrics = fund_map.get(ticker, {}).get("metrics", {})
        growth_data.append({
            "Ticker": ticker,
            "Revenue Growth %": metrics.get("revenue_growth"),
            "Earnings Growth %": metrics.get("earnings_growth"),
            "R&D % Revenue": metrics.get("rd_percent"),
            "FCF Yield %": metrics.get("fcf_yield"),
        })
    growth_df = pd.DataFrame(growth_data)
    if not growth_df.empty:
        _add_dataframe_to_sheet(ws, growth_df)
    ws.freeze_panes = "A2"

    # ========== Sheet 6: Recommendation ==========
    ws = wb.create_sheet("Recommendation", 5)
    rec_data = []
    for _, row in shortlist_df.iterrows():
        ticker = row["Ticker"]
        rec = fund_map.get(ticker, {}).get("recommendation", {})
        rec_data.append({
            "Ticker": ticker,
            "Action": rec.get("action", "HOLD"),
            "Conviction": rec.get("conviction", 0),
            "Entry Low": rec.get("entry_zone", (None, None))[0],
            "Entry High": rec.get("entry_zone", (None, None))[1],
            "Stop Loss": rec.get("stop_loss"),
            "Target 1": rec.get("targets", {}).get("target1"),
            "Target 2": rec.get("targets", {}).get("target2"),
            "Bull Case": " | ".join(rec.get("bull_case", [])[:3]),
            "Bear Case": " | ".join(rec.get("bear_case", [])[:2]),
        })
    rec_df = pd.DataFrame(rec_data)
    if not rec_df.empty:
        _add_dataframe_to_sheet(ws, rec_df)
    ws.freeze_panes = "A2"

    # ========== Per-Stock Deep Dives (Sheets 7+) ==========
    max_sheets = min(25, len(shortlist_df))
    for idx, (_, row) in enumerate(shortlist_df.head(max_sheets).iterrows(), 1):
        ticker = row["Ticker"]
        metrics = fund_map.get(ticker, {}).get("metrics", {})
        technicals = fund_map.get(ticker, {}).get("technicals", {})
        rec = fund_map.get(ticker, {}).get("recommendation", {})

        ws = wb.create_sheet(f"{ticker}", 6 + idx - 1)

        # Header with stock info
        ws["A1"] = f"{ticker} - Deep Dive Analysis"
        ws["A1"].font = Font(size=14, bold=True, color="1F4E78")

        # Fundamentals section
        row_num = 3
        ws[f"A{row_num}"] = "FUNDAMENTALS"
        ws[f"A{row_num}"].font = Font(bold=True, size=11)

        fundamentals = [
            ("Price", metrics.get("price")),
            ("Market Cap", metrics.get("market_cap")),
            ("Sector", metrics.get("sector")),
            ("Industry", metrics.get("industry")),
            ("P/E", metrics.get("pe_trailing")),
            ("P/B", metrics.get("pb")),
            ("ROE %", metrics.get("roe")),
            ("ROCE %", metrics.get("roce")),
            ("Piotroski", metrics.get("piotroski_score")),
            ("FCF Yield %", metrics.get("fcf_yield")),
            ("EV/FCF", metrics.get("ev_fcf")),
            ("D/E", metrics.get("debt_to_equity")),
            ("R&D % Revenue", metrics.get("rd_percent")),
        ]

        for key, value in fundamentals:
            row_num += 1
            ws[f"A{row_num}"] = key
            ws[f"B{row_num}"] = value

        # Technicals section
        row_num += 2
        ws[f"A{row_num}"] = "TECHNICALS"
        ws[f"A{row_num}"].font = Font(bold=True, size=11)

        technicals_list = [
            ("Price $", technicals.get("Price $")),
            ("RSI", technicals.get("RSI")),
            ("ADX", technicals.get("ADX")),
            ("Stage", technicals.get("Stage")),
            ("Vol Ratio", technicals.get("Vol Ratio")),
            ("Target $", technicals.get("Target $")),
            ("Upside %", technicals.get("Upside %")),
        ]

        for key, value in technicals_list:
            row_num += 1
            ws[f"A{row_num}"] = key
            ws[f"B{row_num}"] = value

        # Recommendation section
        row_num += 2
        ws[f"A{row_num}"] = "RECOMMENDATION"
        ws[f"A{row_num}"].font = Font(bold=True, size=11)

        row_num += 1
        ws[f"A{row_num}"] = "Action"
        ws[f"B{row_num}"] = rec.get("action", "HOLD")

        row_num += 1
        ws[f"A{row_num}"] = "Conviction"
        ws[f"B{row_num}"] = rec.get("conviction", 0)

        row_num += 1
        ws[f"A{row_num}"] = "Entry Zone"
        entry = rec.get("entry_zone", (None, None))
        ws[f"B{row_num}"] = f"${entry[0]:.2f} - ${entry[1]:.2f}" if entry[0] else "N/A"

        row_num += 1
        ws[f"A{row_num}"] = "Stop Loss"
        ws[f"B{row_num}"] = f"${rec.get('stop_loss', 0):.2f}" if rec.get("stop_loss") else "N/A"

        row_num += 1
        ws[f"A{row_num}"] = "Bull Case"
        ws[f"B{row_num}"] = "\n".join(rec.get("bull_case", []))

        row_num += 1
        ws[f"A{row_num}"] = "Bear Case"
        ws[f"B{row_num}"] = "\n".join(rec.get("bear_case", []))

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 50

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
